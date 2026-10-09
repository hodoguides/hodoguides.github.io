import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
d = json.load(open(src, encoding="utf-8"))
CAT_FR = {"food": "Manger", "cafe": "Café", "see": "À voir", "photo": "Spot photo", "activity": "Activité", "sleep": "Dormir", "hidden": "Pépite cachée"}

F = "Arial"
ink = "3B3A5E"
hdr_font = Font(name=F, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=ink)
edit_fill = PatternFill("solid", fgColor="FFF3D6")      # cells to fill in
body = Font(name=F, size=10)
thin = Side(style="thin", color="E6DDD0")
grid = Border(bottom=thin)

wb = Workbook()

# ---------- Mode d'emploi ----------
ws = wb.active
ws.title = "Mode d'emploi"
lines = [
    ("HodoGuides · tri de ta liste Google Maps « Japon »", Font(name=F, bold=True, size=15, color=ink)),
    ("", None),
    ("Remplis seulement les cases jaunes de l'onglet « Lieux à trier » :", Font(name=F, bold=True)),
    ("1. À garder ?  →  Oui = le lieu va sur la carte · Non = je l'enlève", body),
    ("2. J'y suis allée ?  →  Oui = testé · Non = sur ma liste « à tester »", body),
    ("3. Mon verdict (seulement si tu y es allée)  →  Coup de cœur · Recommandé · Si tu as le temps", body),
    ("4. Catégorie  →  déjà devinée, corrige-la si besoin", body),
    ("5. Mon avis en une phrase  →  facultatif, tu pourras l'écrire plus tard", body),
    ("", None),
    ("Astuce : utilise les filtres de la ligne de titre (par ville, par catégorie) pour avancer région par région.", body),
    ("Les cases laissées vides sont ignorées : pas besoin de tout faire d'un coup.", body),
    ("Les positions sont approximatives (à quelques km) : je les préciserai pour les lieux que tu gardes.", body),
    ("", None),
    ("Exemple de ligne remplie :", Font(name=F, bold=True)),
]
for i, (text, font) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=text)
    if font: c.font = font
heads = ["À garder ?", "J'y suis allée ?", "Mon verdict", "Catégorie", "Nom", "Ville (approx.)", "Mon avis en une phrase"]
ex = ["Oui", "Oui", "Coup de cœur", "À voir", "Fushimi Inari-taisha", "Kyoto", "Arrive avant 8h, la foule disparaît après la moitié."]
r0 = len(lines) + 1
for j, (h, v) in enumerate(zip(heads, ex), 1):
    ws.cell(row=r0, column=j, value=h).font = hdr_font
    ws.cell(row=r0, column=j).fill = hdr_fill
    c = ws.cell(row=r0 + 1, column=j, value=v); c.font = body
    if j in (1, 2, 3, 4, 7): c.fill = edit_fill
ws.cell(row=r0 + 3, column=1, value="Où tu en es :").font = Font(name=F, bold=True)
n = len(d["keep"])
last = n + 1
stats = [
    ("Lieux dans la liste", f"=COUNTA('Lieux à trier'!E2:E{last})"),
    ("Déjà triés (Oui ou Non)", f"=COUNTIF('Lieux à trier'!A2:A{last},\"Oui\")+COUNTIF('Lieux à trier'!A2:A{last},\"Non\")"),
    ("Gardés pour la carte", f"=COUNTIF('Lieux à trier'!A2:A{last},\"Oui\")"),
    ("Dont testés par toi", f"=COUNTIFS('Lieux à trier'!A2:A{last},\"Oui\",'Lieux à trier'!B2:B{last},\"Oui\")"),
    ("Dont coups de cœur", f"=COUNTIFS('Lieux à trier'!A2:A{last},\"Oui\",'Lieux à trier'!C2:C{last},\"Coup de cœur\")"),
]
for k, (lab, f) in enumerate(stats):
    ws.cell(row=r0 + 4 + k, column=1, value=lab).font = body
    c = ws.cell(row=r0 + 4 + k, column=2, value=f); c.font = Font(name=F, bold=True)
ws.column_dimensions["A"].width = 26
for col, w in zip("BCDEFG", (16, 16, 14, 26, 16, 48)): ws.column_dimensions[col].width = w

# ---------- Lieux à trier ----------
ws = wb.create_sheet("Lieux à trier")
cols = ["À garder ?", "J'y suis allée ?", "Mon verdict", "Catégorie", "Nom", "Ville (approx.)", "Mon avis en une phrase", "Lien Google Maps", "lat", "lng"]
widths = [12, 15, 17, 15, 42, 18, 46, 18, 9, 9]
for j, (h, w) in enumerate(zip(cols, widths), 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(vertical="center")
    ws.column_dimensions[get_column_letter(j)].width = w
rows = sorted(d["keep"], key=lambda k: (k["city"].replace("près de ", "~"), k["category"], k["name"]))
for i, k in enumerate(rows, 2):
    vals = [None, None, None, CAT_FR[k["category"]], k["name"], k["city"], k["note"] or None, "Ouvrir", k["lat"], k["lng"]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = body; c.border = grid
        if j in (1, 2, 3, 4, 7): c.fill = edit_fill
    link = ws.cell(row=i, column=8); link.hyperlink = k["url"]; link.font = Font(name=F, size=10, color="2F6FB5", underline="single")
for col, opts in (("A", "Oui,Non"), ("B", "Oui,Non"), ("C", "Coup de cœur,Recommandé,Si tu as le temps"), ("D", ",".join(CAT_FR.values()))):
    dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
ws.freeze_panes = "F2"
ws.auto_filter.ref = f"A1:J{last}"
ws.column_dimensions["I"].hidden = True
ws.column_dimensions["J"].hidden = True

# ---------- Villes & régions ----------
ws = wb.create_sheet("Villes & régions")
for j, (h, w) in enumerate(zip(["Faire une page ?", "Nom", "Lien Google Maps"], (16, 36, 18)), 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
    ws.column_dimensions[get_column_letter(j)].width = w
ws.cell(row=1, column=5, value="Ce sont des villes ou régions entières, pas des lieux précis : elles pourront devenir des sections de la page Japon.").font = Font(name=F, italic=True, size=10)
for i, a in enumerate(sorted(d["areas"], key=lambda a: a["name"]), 2):
    ws.cell(row=i, column=1).fill = edit_fill
    ws.cell(row=i, column=2, value=a["name"]).font = body
    l = ws.cell(row=i, column=3, value="Ouvrir"); l.hyperlink = a["url"]; l.font = Font(name=F, size=10, color="2F6FB5", underline="single")
dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"A2:A{len(d['areas'])+1}")

# ---------- Retirés ----------
ws = wb.create_sheet("Retirés (à vérifier)")
for j, (h, w) in enumerate(zip(["Remettre ?", "Nom", "Pourquoi"], (12, 46, 30)), 1):
    c = ws.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill
    ws.column_dimensions[get_column_letter(j)].width = w
WHY = {"tri manuel": "adresse, salon, service…", "pas de position": "repère sans nom ni position"}
items = [(x["name"], WHY.get(x["why"], "hôtel, gare, parking, supérette…")) for x in d["dropped"]] + [(x["name"], "hors du Japon") for x in d["outside"]]
for i, (name, why) in enumerate(items, 2):
    ws.cell(row=i, column=1).fill = edit_fill
    ws.cell(row=i, column=2, value=name).font = body
    ws.cell(row=i, column=3, value=why).font = body
dv = DataValidation(type="list", formula1='"Oui"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"A2:A{len(items)+1}")

wb.active = 1
wb.save(out)
print(out, n, len(d["areas"]), len(items))
